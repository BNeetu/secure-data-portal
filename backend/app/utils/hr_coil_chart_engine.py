"""
HR Coil SUMMARY Sheet Chart Engine
====================================
Detects every pivot-table block in the SUMMARY sheet, extracts
(title, label-column, %WRP-column) for each, generates the best chart,
and inserts it neatly in that table's section.

Works for any monthly HR Coil report — no hardcoded cell ranges.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.chart import BarChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter

# ── Modern steel-industry colour palette ─────────────────────────────────────
PALETTE = [
    "4472C4","ED7D31","A9D18E","FF0000","FFC000",
    "5B9BD5","70AD47","255E91","9DC3E6","F4B183",
    "C9C9C9","7F7F7F","4BACC6","8064A2","9BBB59",
]

SKIP_LABELS = {
    "grand total","total","sum of qty(mt)","sum of qty","nos coils",
    "count of batch","%wrp","% wrp","row labels","rolled grade","grades",
}

def _val(grid, r, c): return grid.get((r, c))

def _is_pct_header(v):
    if v is None: return False
    s = str(v).strip().upper()
    return s in ("%WRP","% WRP","%","PERCENTAGE","PCT","% SHARE","PERC")

def _is_qty_header(v):
    if v is None: return False
    s = str(v).strip().upper()
    return any(k in s for k in ("QTY","QUANTITY","SUM OF","WEIGHT","MT","NOS","COUNT","BATCH"))

def _is_title(v):
    if v is None: return False
    if isinstance(v, (int,float)): return False
    s = str(v).strip()
    if not s: return False
    return s.lower() not in SKIP_LABELS and not _is_pct_header(v) and not _is_qty_header(v)

@dataclass
class TableBlock:
    title:      str
    header_row: int
    data_start: int
    data_end:   int
    label_col:  int
    pct_col:    int

    @property
    def n_rows(self): return max(0, self.data_end - self.data_start + 1)


def detect_tables(ws) -> List[TableBlock]:
    """Find every table in the sheet dynamically."""
    grid: Dict[Tuple[int,int], Any] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                grid[(cell.row, cell.column)] = cell.value

    max_row = ws.max_row

    # Find every %WRP header cell
    pct_cells = [(r,c) for (r,c),v in grid.items() if _is_pct_header(v)]
    pct_cells.sort()

    tables: List[TableBlock] = []
    seen = set()

    for (hdr_row, pct_col) in pct_cells:
        key = (hdr_row, pct_col)
        if key in seen: continue
        seen.add(key)

        # Find label column: walk left from pct_col, skip qty-header cols
        label_col = None
        for c in range(pct_col - 1, 0, -1):
            v = _val(grid, hdr_row, c)
            if v is not None and not _is_pct_header(v) and not _is_qty_header(v):
                label_col = c
                break

        if label_col is None:
            continue

        # Find the table title:
        # Priority 1: value in label_col at hdr_row (same row, e.g. "QC STATUS")
        # Priority 2: scan rows above for the nearest title-like string
        title = ""
        candidate = _val(grid, hdr_row, label_col)
        if _is_title(candidate):
            title = str(candidate).strip()
        else:
            for tr in range(hdr_row - 1, max(0, hdr_row - 5), -1):
                for c in range(label_col, pct_col + 1):
                    v = _val(grid, tr, c)
                    if _is_title(v):
                        title = str(v).strip()
                        break
                if title:
                    break

        if not title:
            title = f"Table (row {hdr_row})"

        # Collect data rows below header
        data_rows = []
        r = hdr_row + 1
        while r <= max_row:
            lv = _val(grid, r, label_col)
            pv = _val(grid, r, pct_col)
            # Stop if both cols empty (gap between tables)
            if lv is None and pv is None:
                break
            if lv is not None:
                lstr = str(lv).strip().lower()
                if lstr == "grand total":
                    break
                if lstr in SKIP_LABELS:
                    r += 1
                    continue
                # Must have a numeric % value to be a real data row
                if pv is not None and isinstance(pv, (int, float)):
                    data_rows.append(r)
            r += 1

        if len(data_rows) < 1:
            continue

        tables.append(TableBlock(
            title      = title,
            header_row = hdr_row,
            data_start = data_rows[0],
            data_end   = data_rows[-1],
            label_col  = label_col,
            pct_col    = pct_col,
        ))

    return tables


def _best_chart_type(n: int) -> str:
    if n <= 3:  return "pie"
    if n <= 6:  return "doughnut"
    if n <= 15: return "bar_vertical"
    return "bar_horizontal"


def _make_chart(ws, table: TableBlock) -> Optional[object]:
    n = table.n_rows
    if n < 1: return None

    chart_type = _best_chart_type(n)
    height = max(8.0, min(18.0, 5.0 + n * 0.9))
    width  = 14.0

    labels_ref = Reference(ws,
        min_col=table.label_col, max_col=table.label_col,
        min_row=table.data_start, max_row=table.data_end)

    values_ref = Reference(ws,
        min_col=table.pct_col, max_col=table.pct_col,
        min_row=table.data_start, max_row=table.data_end)

    if chart_type == "pie":
        chart = PieChart()
        chart.style = 26
    elif chart_type == "doughnut":
        chart = DoughnutChart()
        chart.style = 26
        chart.holeSize = 40
    else:
        chart = BarChart()
        chart.style = 26
        chart.type    = "bar" if chart_type == "bar_horizontal" else "col"
        chart.barDir  = "bar" if chart_type == "bar_horizontal" else "col"
        chart.grouping = "clustered"
        chart.overlap  = 0

    series = openpyxl.chart.Series(values_ref, title=table.title)
    chart.series.append(series)
    chart.set_categories(labels_ref)

    # Data labels
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showPercent = chart_type in ("pie","doughnut")
    chart.dataLabels.showCatName = chart_type in ("pie","doughnut")
    chart.dataLabels.showLegendKey = False

    # Color each slice/bar individually
    for i in range(n):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
        chart.series[0].dPt.append(pt)

    chart.title   = table.title
    chart.width   = width
    chart.height  = height

    if hasattr(chart, 'legend') and chart.legend:
        chart.legend.position = "b"

    return chart


def generate_and_embed_charts(raw_bytes: bytes) -> bytes:
    """
    Main entry: detect all tables in SUMMARY, generate one chart per table,
    embed each chart at anchor = (label_col+1, header_row), return new bytes.
    """
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes))

    # Find SUMMARY sheet (case-insensitive)
    ws = None
    for name in wb.sheetnames:
        if "summary" in name.lower():
            ws = wb[name]
            break
    if ws is None:
        return raw_bytes

    tables = detect_tables(ws)
    if not tables:
        return raw_bytes

    placed = []  # track anchor positions to avoid exact overlaps

    for table in tables:
        chart = _make_chart(ws, table)
        if chart is None:
            continue

        # Anchor: one column right of label_col, at header_row
        anchor_col = table.label_col + 1
        anchor_row = table.header_row
        anchor_str = f"{get_column_letter(anchor_col)}{anchor_row}"

        if anchor_str in placed:
            anchor_row += 1
            anchor_str = f"{get_column_letter(anchor_col)}{anchor_row}"

        placed.append(anchor_str)
        ws.add_chart(chart, anchor_str)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
